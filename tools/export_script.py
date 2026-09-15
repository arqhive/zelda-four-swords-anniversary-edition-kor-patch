"""Export translation sheet: JP (source) + EN (reference) + empty KO, per kmsg entry.

Usage: python tools/export_script.py --rom path/to/00000000 [--out-dir work/translation]

Outputs (contain original game text - do NOT commit):
  script.json  - JP/EN markup + plain text per entry
  script.xlsx  - review sheet (if openpyxl is available)
"""
import argparse, json, os, sys
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ndsrom import Rom
from kmsg import Kmsg, to_markup, from_markup, plain_text, build_string

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ap = argparse.ArgumentParser()
ap.add_argument("--rom", required=True)
ap.add_argument("--out-dir", default=os.path.join(REPO, "work", "translation"))
args = ap.parse_args()
OUT = args.out_dir
os.makedirs(OUT, exist_ok=True)

rom = Rom(args.rom)
raw_all, raw_us = rom.read("all.kmsg"), rom.read("us.kmsg")
ka, ku = Kmsg(raw_all), Kmsg(raw_us)
assert ka.build() == raw_all, "all.kmsg roundtrip mismatch"
assert ku.build() == raw_us, "us.kmsg roundtrip mismatch"

rows = []
for i, ((_, la), (_, lu)) in enumerate(zip(ka.entries, ku.entries)):
    jp, en = la[0], lu[1]
    # markup must roundtrip for both languages
    assert from_markup(to_markup(jp)) == jp, "JP markup roundtrip %d" % i
    assert from_markup(to_markup(en)) == en, "EN markup roundtrip %d" % i
    assert build_string(from_markup(to_markup(en))) == build_string(en)
    rows.append({"id": i, "jp": to_markup(jp), "en": to_markup(en), "ko": "",
                 "jp_plain": plain_text(jp), "en_plain": plain_text(en)})

with open(os.path.join(OUT, "script.json"), "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False, indent=1)
jp_chars = sum(len(r["jp_plain"].replace("\n", "")) for r in rows)
print("exported %d entries, JP plain chars %d -> %s" % (len(rows), jp_chars, os.path.join(OUT, "script.json")))

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook(); ws = wb.active; ws.title = "script"
    ws.append(["id", "일본어 (원문)", "영어 (참고)", "한국어 (번역)", "일본어 markup", "영어 markup"])
    for r in rows:
        ws.append([r["id"], r["jp_plain"], r["en_plain"], r["ko"], r["jp"], r["en"]])
    for col, w in zip("ABCDEF", (6, 50, 50, 50, 60, 60)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2"
    wb.save(os.path.join(OUT, "script.xlsx"))
    print("xlsx written")
except ImportError:
    print("openpyxl not installed; skipped xlsx")
